import csv
import io
import re
import sys
import urllib.request
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
LEVELS_URL = "https://www.ine.es/daco/daco42/clasificaciones/cned14/CNED14_A_estr.xls"
FIELDS_URL = "https://www.ine.es/daco/daco42/clasificaciones/cned14/CNED_F_estr.xlsx"
USER_AGENT = "Mozilla/5.0 (compatible; datos-referencia; +https://github.com/cristianemek/datos-referencia)"
HEADER = ["code", "level", "name", "parent_code"]
LEVEL_PATTERNS_A = [(1, r"[A-Z]"), (2, r"\d"), (3, r"[A-Z]{2}"), (4, r"\d{2}")]
LEVEL_PATTERNS_F = [(1, r"\d{2}"), (2, r"\d{3}"), (3, r"\d{4}")]
SPREADSHEET_ERRORS = {"#REF!", "#¡REF!", "#N/A", "#VALUE!", "#¡VALOR!", "#NAME?", "#¿NOMBRE?"}


def fetch(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})

    with urllib.request.urlopen(request, timeout=120) as response:
        return response.read()


def open_sheet(url: str):
    content = fetch(url)

    if not content.startswith(b"PK"):
        raise RuntimeError(f"{url} no es un libro OOXML")

    return openpyxl.load_workbook(io.BytesIO(content), data_only=True).worksheets[0]


def clean(value) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()

    return "" if text in SPREADSHEET_ERRORS else text


def level_of(code: str, patterns: list) -> int:
    return next((level for level, pattern in patterns if re.fullmatch(pattern, code)), 0)


def read_levels() -> list:
    sheet = open_sheet(LEVELS_URL)
    rows = []
    last = {}

    for code_cell, name_cell, *_ in sheet.iter_rows(values_only=True):
        code, name = clean(code_cell), clean(name_cell)
        level = level_of(code, LEVEL_PATTERNS_A)

        if not level or not name:
            continue

        parent = last.get(level - 1, "")

        if level > 1 and not parent:
            raise RuntimeError(f"CNED-A {code}: sin epígrafe superior")

        last = {**{k: v for k, v in last.items() if k < level}, level: code}
        rows.append({"code": code, "level": level, "name": name, "parent_code": parent})

    return rows


def validate_levels(rows: list) -> None:
    by_code = {row["code"]: row for row in rows}

    if len(by_code) != len(rows):
        raise RuntimeError("Hay códigos CNED-A duplicados")

    for row in rows:
        parent = by_code.get(row["parent_code"])

        if row["level"] in (3, 4) and row["code"][0] != parent["parent_code"]:
            raise RuntimeError(f"CNED-A {row['code']}: su primer carácter no es el del epígrafe de dos niveles más arriba")


def split_entry(text: str) -> tuple:
    match = re.fullmatch(r"(\d{2,4}) (.+)", text)

    if not match:
        raise RuntimeError(f"CNED-F: celda sin código reconocible: {text!r}")

    return match.group(1), match.group(2)


def read_fields() -> list:
    sheet = open_sheet(FIELDS_URL)
    rows = []
    last = {}

    for values in sheet.iter_rows(min_row=3, max_col=3, values_only=True):
        for column, value in enumerate(values, start=1):
            text = clean(value)

            if not text:
                continue

            code, name = split_entry(text)
            level = level_of(code, LEVEL_PATTERNS_F)

            if level != column:
                raise RuntimeError(f"CNED-F {code}: aparece en la columna {column}")

            parent = code[:-1] if level == 3 else code[:2] if level == 2 else ""

            if level > 1 and last.get(level - 1) != parent:
                raise RuntimeError(f"CNED-F {code}: el sector superior de la hoja no es {parent}")

            last = {**{k: v for k, v in last.items() if k < level}, level: code}
            rows.append({"code": code, "level": level, "name": name, "parent_code": parent})

    if len({row["code"] for row in rows}) != len(rows):
        raise RuntimeError("Hay códigos CNED-F duplicados")

    return rows


def write_csv(path: Path, rows: list) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(HEADER)
    writer.writerows([row[key] for key in HEADER] for row in rows)
    path.write_text(buffer.getvalue(), encoding="utf-8")


def summary(rows: list, depth: int) -> list:
    return [sum(1 for row in rows if row["level"] == level) for level in range(1, depth + 1)]


def main() -> None:
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / "catalogos"
    levels = read_levels()
    validate_levels(levels)
    fields = read_fields()
    write_csv(folder / "cned_2014_niveles.csv", levels)
    write_csv(folder / "cned_2014_sectores.csv", fields)
    print(f"CNED-A: {len(levels)} epígrafes {summary(levels, 4)}; CNED-F: {len(fields)} epígrafes {summary(fields, 3)}")


if __name__ == "__main__":
    sys.exit(main())
