import re
from collections import Counter, defaultdict
from decimal import ROUND_HALF_UP, Decimal

import openpyxl

from comun import BUILD, REPO, fail_if, fetch, fold, read_csv, write_csv

LAU_URL = "https://ec.europa.eu/eurostat/documents/345175/501971/EU-27-LAU-2025-NUTS-2024.xlsx"
CENSO_URL = "https://www.mjusticia.gob.es/es/JusticiaEspana/OrganizacionJusticia/Documents/Censo%20Judicial.xlsx"
MUNICIPIOS = REPO / "direcciones" / "municipios_es.csv"
OUTPUT = BUILD / "municipios_columnas.csv"
PARTIDOS = REPO / "direcciones" / "partidos_judiciales_es.csv"

HEADER = [
    "municipality_code", "area_km2", "degurba", "coastal", "city_code", "city_name",
    "fua_code", "fua_name", "judicial_district_code", "judicial_district",
]
LAU_COLUMNS = {
    "nuts3": "NUTS3", "code": "LAU CODE", "area": "TOTAL AREA (m2)", "degurba": "DEGURBA",
    "coastal": "COASTAL AREA", "city_code": "CITY ID", "city_name": "CITY NAME",
    "fua_code": "FUA ID", "fua_name": "FUA_NAME",
}
CENSO_COLUMNS = {
    "province": "CPRO", "seat_flag": "CPJ", "number": "PJ", "district": "CODM PJ",
    "district_name": "PARTIDO", "code": "CODM", "name": "LOCALIDAD",
}


def text(value) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def km2(value) -> str:
    return str((Decimal(str(value)) / Decimal(1_000_000)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def column_index(header: tuple, wanted: dict) -> dict:
    names = [text(h) for h in header]
    missing = [v for v in wanted.values() if v not in names]
    fail_if(missing, "Faltan columnas en el fichero de origen")
    return {key: names.index(value) for key, value in wanted.items()}


def load_lau() -> dict:
    sheet = openpyxl.load_workbook(fetch(LAU_URL, "EU-27-LAU-2025-NUTS-2024.xlsx"), read_only=True)["ES"]
    rows = sheet.iter_rows(values_only=True)
    index = column_index(next(rows), LAU_COLUMNS)
    lau = {}
    for row in rows:
        record = {key: row[i] for key, i in index.items()}
        code = text(record["code"])
        if not code:
            continue
        fail_if([code] if code in lau else [], "Código LAU repetido")
        lau[code] = {
            "nuts3": text(record["nuts3"]),
            "area_km2": km2(record["area"]) if record["area"] not in (None, "") else "",
            "degurba": text(record["degurba"]),
            "coastal": text(record["coastal"]),
            "city_code": text(record["city_code"]),
            "city_name": text(record["city_name"]),
            "fua_code": text(record["fua_code"]),
            "fua_name": text(record["fua_name"]),
        }
    return lau


def check_lau(lau: dict, municipios: list[dict]) -> None:
    codes = {m["municipality_code"] for m in municipios}
    fail_if(sorted(codes - set(lau)), "Municipios sin fila LAU")
    fail_if(sorted(set(lau) - codes), "Filas LAU sin municipio INE")
    fail_if([m["municipality_code"] for m in municipios if lau[m["municipality_code"]]["nuts3"] != m["nuts3"]],
            "NUTS3 distinto entre LAU y municipios_es.csv")
    fail_if([c for c, r in lau.items() if r["degurba"] not in {"1", "2", "3"}], "DEGURBA fuera de rango")
    fail_if([c for c, r in lau.items() if r["coastal"] not in {"0", "1"}], "COASTAL AREA fuera de rango")
    fail_if([c for c, r in lau.items() if not r["area_km2"] or Decimal(r["area_km2"]) <= 0], "Superficie vacía o nula")
    fail_if([c for c, r in lau.items() if r["city_code"] and not r["fua_code"]], "Ciudad sin FUA")
    for key in ("city", "fua"):
        names = defaultdict(set)
        for r in lau.values():
            if r[f"{key}_code"]:
                names[r[f"{key}_code"]].add(r[f"{key}_name"])
        fail_if([f"{c}: {sorted(n)}" for c, n in names.items() if len(n) != 1 or "" in n], f"Códigos {key} con varios nombres")


def load_censo() -> tuple[dict, str]:
    sheet = openpyxl.load_workbook(fetch(CENSO_URL, "Censo Judicial.xlsx"), read_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    index = column_index(next(rows), CENSO_COLUMNS)
    censo, updated = {}, ""
    for row in rows:
        record = {key: text(row[i]) for key, i in index.items()}
        first = text(row[0])
        if first.lower().startswith("actualizado"):
            updated = first
        if not record["code"]:
            continue
        fail_if([record["code"]] if record["code"] in censo else [], "Municipio repetido en el Censo Judicial")
        censo[record["code"]] = record
    return censo, updated


def build_partidos(censo: dict, names: dict) -> list[dict]:
    groups = defaultdict(list)
    for record in censo.values():
        groups[record["district"]].append(record)
    partidos, problems, variants = [], [], []
    for code, members in sorted(groups.items()):
        seats = [m for m in members if m["seat_flag"].upper() == "SI"]
        numbers = {m["number"] for m in members}
        provinces = {m["province"].zfill(2) for m in members}
        if len(seats) != 1 or seats[0]["code"] != code:
            problems.append(f"{code}: cabecera {[s['code'] for s in seats]}")
        if len(numbers) != 1 or len(provinces) != 1 or code[:2] not in provinces:
            problems.append(f"{code}: número {numbers} o provincia {provinces} incoherentes")
        labels = Counter(m["district_name"] for m in members)
        if len(labels) > 1 or fold(next(iter(labels))) != fold(censo[code]["name"]):
            variants.append(f"{code} {dict(labels)} -> {names[code]}")
        partidos.append({"code": code, "name": names[code], "province_code": code[:2], "number": numbers.pop()})
    fail_if(problems, "Partidos judiciales incoherentes")
    for line in variants:
        print(f"  aviso, nombre de partido normalizado: {line}")
    return partidos


def main() -> None:
    municipios = read_csv(MUNICIPIOS)
    names = {m["municipality_code"]: m["municipality"] for m in municipios}
    lau = load_lau()
    check_lau(lau, municipios)
    censo, updated = load_censo()
    fail_if(sorted(set(names) - set(censo)), "Municipios sin partido judicial")
    fail_if(sorted(set(censo) - set(names)), "Municipios del Censo Judicial que no están en el INE")
    partidos = build_partidos(censo, names)
    rows = []
    for m in municipios:
        code = m["municipality_code"]
        district = censo[code]["district"]
        fields = {k: lau[code][k] for k in HEADER[1:8]}
        rows.append({"municipality_code": code, **fields,
                     "judicial_district_code": district, "judicial_district": names[district]})
    write_csv(OUTPUT, HEADER, rows)
    write_csv(PARTIDOS, ["code", "name", "province_code", "number"], partidos)
    print(f"Censo Judicial: {updated or 'sin fecha de actualización'}")


if __name__ == "__main__":
    main()
