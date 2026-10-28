import argparse
import csv
import datetime
import io
import re
import sys
import unicodedata
import urllib.request
from collections import defaultdict
from pathlib import Path
from typing import NamedTuple, Optional

import openpyxl

REPO = Path(__file__).resolve().parents[2]
SOURCE_URL = "https://www.bde.es/f/webbe/SGE/regis/ficheros/es/Registro_Entidades.xlsx"
USER_AGENT = "datos-referencia (https://github.com/cristianemek/datos-referencia)"
SHEET = "DATOS GENERALES"
HEADER_MARK = "CÓDIGO BE"
BANK_CODE = re.compile(r"\d{4}")
AS_OF = re.compile(r"FECHA:\s*(\d{2}/\d{2}/\d{4})")

FIELDS = {
    "code": "CÓDIGO BE",
    "name": "NOMBRE",
    "nif": "DOCUMENTO IDENTIFICATIVO",
    "type_code": "CODIGO TIPO ENTIDAD",
    "type_name": "TIPO ENTIDAD",
    "lei": "CÓDIGO LEI",
    "registered_on": "FECHA ALTA",
    "deregistered_on": "FECHA BAJA",
    "reason": "MOTIVO BAJA",
    "successor_code": "SUCESORAS",
    "web": "PÁGINA WEB",
}

ENTITY_TYPES = {
    "TEBP": "banco",
    "TEBL": "banco_local",
    "TEBN": "banco_nacional",
    "TEBR": "banco_regional",
    "TEBEX": "banco_extranjero",
    "TECA": "caja_ahorros",
    "TECC": "cooperativa_credito",
    "TECCC": "cooperativa_credito_calificada",
    "TEOCC": "cooperativa_credito_otra",
    "TECRC": "caja_rural_calificada",
    "TECRNC": "caja_rural_no_calificada",
    "TECO": "credito_oficial",
    "TEECCP": "entidad_ley_25_1991",
    "TESECC": "sucursal_ue",
    "TESECE": "sucursal_extracomunitaria",
    "TEOR": "oficina_representacion",
    "TEEFC": "efc",
    "TEEFEP": "efc_entidad_pago",
    "TEEFPH": "efc_entidad_pago_hibrida",
    "TEEP": "entidad_pago",
    "TEEPEX": "entidad_pago_exenta",
    "TEEPH": "entidad_pago_hibrida",
    "TEEDE": "entidad_dinero_electronico",
    "TEPSIC": "proveedor_informacion_cuentas",
    "TESEPC": "sucursal_ue_entidad_pago",
    "TESEDC": "sucursal_ue_dinero_electronico",
    "TESEFC": "sucursal_ue_filial_financiera",
    "TEEF": "entidad_financiacion",
    "TESF": "entidad_factoring",
    "TESAF": "sociedad_arrendamiento_financiero",
    "TESCH": "sociedad_credito_hipotecario",
    "TESMMD": "sociedad_mediadora_mercado_dinero",
    "TESDEC": "sociedad_dominante_entidad_credito",
    "TESFC": "sociedad_financiera_cartera",
    "TESFMC": "sociedad_financiera_mixta_cartera",
    "TESGR": "sociedad_garantia_reciproca",
    "TESR": "sociedad_reafianzamiento",
    "TEST": "sociedad_tasacion",
    "TEECVM": "compraventa_moneda",
    "TEECVT": "compraventa_moneda_transferencias",
    "CRIPTO": "proveedor_moneda_virtual",
}

CREDIT_TYPES = ("TEBP", "TECA", "TECC", "TECO", "TESECC", "TESECE")
PAYMENT_TYPES = ("TEEP", "TEEPEX", "TEEDE", "TEPSIC", "TEEFC", "TEEFEP", "TESEPC", "TESEDC", "TESEFC")


class Renumbering(NamedTuple):
    code: str
    nif: str
    until: datetime.date
    source: str


RENUMBERED = {
    "0224": Renumbering("0073", "A28122570", datetime.date(2026, 5, 4), "BOE-A-2026-11394"),
}

ENRICHMENT_HEADER = ["bank_code", "nif", "lei", "entity_type", "active", "deregistered_on", "successor_code"]
HISTORICAL_HEADER = [
    "bank_code", "name", "entity_type", "registered_on", "deregistered_on", "reason", "successor_code", "successor_name",
]
PAYMENTS_HEADER = ["bank_code", "name", "nif", "lei", "entity_type", "registered_on"]
TYPES_HEADER = ["entity_type", "bde_type_code", "bde_type_name"]
NEW_BANKS_HEADER = ["bank_code", "name", "nif", "lei", "entity_type", "web"]


class Entry(NamedTuple):
    code: str
    name: str
    nif: str
    lei: str
    type_code: str
    type_name: str
    registered_on: datetime.date
    deregistered_on: Optional[datetime.date]
    reason: str
    successor_code: str
    web: str


class Index(NamedTuple):
    by_code: dict
    by_identity: dict


def fetch(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})

    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            return response.read()
    except OSError as error:
        raise RuntimeError(f"No se pudo descargar {url}: {error}") from error


def text(value) -> str:
    return "" if value is None else str(value)


def parse_date(value: str) -> Optional[datetime.date]:
    if not value:
        return None

    try:
        return datetime.datetime.strptime(value, "%d/%m/%Y").date()
    except ValueError as error:
        raise ValueError(f"Fecha no válida en el registro: {value!r}") from error


def iso(value: Optional[datetime.date]) -> str:
    return value.isoformat() if value else ""


def header_positions(header: tuple) -> dict:
    labels = [text(cell).strip() for cell in header]
    missing = [label for label in FIELDS.values() if label not in labels]

    if missing:
        raise ValueError(f"Faltan columnas en la hoja {SHEET}: {', '.join(missing)}")

    return {field: labels.index(label) for field, label in FIELDS.items()}


def to_entry(row: tuple, positions: dict) -> Entry:
    values = {field: text(row[position]) for field, position in positions.items()}

    return Entry(
        **{
            **values,
            "registered_on": parse_date(values["registered_on"]),
            "deregistered_on": parse_date(values["deregistered_on"]),
        }
    )


def read_register(data: bytes) -> tuple:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)

    if SHEET not in workbook.sheetnames:
        raise ValueError(f"El fichero no tiene la hoja {SHEET}")

    rows = iter(workbook[SHEET].iter_rows(values_only=True))
    as_of = ""

    for row in rows:
        match = AS_OF.search(text(row[0]))
        as_of = match.group(1) if match else as_of

        if text(row[0]).strip() == HEADER_MARK:
            positions = header_positions(row)
            break
    else:
        raise ValueError(f"No se encuentra la cabecera de la hoja {SHEET}")

    entries = tuple(to_entry(row, positions) for row in rows if row and row[0] is not None)
    return as_of, tuple(entry for entry in entries if BANK_CODE.fullmatch(entry.code))


def identity(entry: Entry) -> str:
    return entry.nif or entry.name


def held_on(entry: Entry, day: datetime.date) -> bool:
    return entry.registered_on <= day and (entry.deregistered_on is None or entry.deregistered_on > day)


def build_index(entries: tuple) -> Index:
    by_code = defaultdict(list)
    by_identity = defaultdict(list)

    for entry in entries:
        by_code[entry.code].append(entry)
        by_identity[identity(entry)].append(entry)

    return Index({k: tuple(v) for k, v in by_code.items()}, {k: tuple(v) for k, v in by_identity.items()})


def entity_type(type_code: str) -> str:
    if type_code not in ENTITY_TYPES:
        raise KeyError(f"Tipo de entidad del BdE sin correspondencia: {type_code!r}")

    return ENTITY_TYPES[type_code]


def active_entry(code: str, index: Index) -> Optional[Entry]:
    active = [entry for entry in index.by_code.get(code, ()) if entry.deregistered_on is None]

    if len(active) > 1:
        raise ValueError(f"El código {code} tiene más de una entidad vigente")

    return active[0] if active else None


def current_or_last(code: str, index: Index) -> Optional[Entry]:
    entries = index.by_code.get(code, ())

    if not entries:
        return None

    return active_entry(code, index) or max(entries, key=lambda entry: (entry.deregistered_on, entry.registered_on))


def renumbered_holder(code: str, index: Index) -> Entry:
    renumbering = RENUMBERED[code]
    holder = active_entry(renumbering.code, index)

    if holder is None or holder.nif != renumbering.nif:
        raise ValueError(f"El cambio de código {code} -> {renumbering.code} ({renumbering.source}) ya no cuadra con el registro")

    return holder


def enrichment_row(code: str, index: Index) -> list:
    entry = current_or_last(code, index)

    if entry is None and code in RENUMBERED:
        holder = renumbered_holder(code, index)
        renumbering = RENUMBERED[code]
        return [code, holder.nif, holder.lei, entity_type(holder.type_code), "0", iso(renumbering.until), renumbering.code]

    if entry is None:
        raise ValueError(f"El código {code} de bancos_es.csv no está en el registro del BdE")

    active = entry.deregistered_on is None
    successor = "" if active else entry.successor_code
    return [code, entry.nif, entry.lei, entity_type(entry.type_code), "1" if active else "0", iso(entry.deregistered_on), successor]


def continues(entry: Entry, index: Index) -> bool:
    return any(
        other is not entry
        and other.code == entry.code
        and other.registered_on <= entry.deregistered_on
        and (other.deregistered_on is None or other.deregistered_on > entry.deregistered_on)
        for other in index.by_identity[identity(entry)]
    )


def single_name(entries) -> str:
    names = {entry.name for entry in entries}
    return names.pop() if len(names) == 1 else ""


def linked_holders(code: str, day: datetime.date, entry: Entry, index: Index) -> tuple:
    identities = {identity(other) for other in index.by_code.get(code, ()) if other is not entry and other.registered_on <= day}
    return tuple(other for key in identities for other in index.by_identity[key] if held_on(other, day))


def renumbered_holders(code: str, day: datetime.date, index: Index) -> tuple:
    renumbering = RENUMBERED.get(code)

    if renumbering is None or day > renumbering.until:
        return ()

    return tuple(other for other in index.by_code.get(renumbering.code, ()) if other.nif == renumbering.nif and held_on(other, day))


def successor_name(entry: Entry, index: Index) -> str:
    code, day = entry.successor_code, entry.deregistered_on

    if not code:
        return ""

    direct = tuple(other for other in index.by_code.get(code, ()) if other is not entry and held_on(other, day))

    if direct:
        return single_name(direct)

    return single_name(linked_holders(code, day, entry, index) or renumbered_holders(code, day, index))


def historical_entries(entries: tuple, index: Index) -> list:
    dropped = [entry for entry in entries if entry.deregistered_on and not continues(entry, index)]
    return sorted(dropped, key=lambda entry: (entry.code, entry.registered_on, entry.deregistered_on, entry.name))


def historical_row(entry: Entry, index: Index) -> list:
    return [
        entry.code, entry.name, entity_type(entry.type_code), iso(entry.registered_on), iso(entry.deregistered_on),
        entry.reason, entry.successor_code, successor_name(entry, index),
    ]


def payment_rows(entries: tuple) -> list:
    active = [entry for entry in entries if entry.deregistered_on is None and entry.type_code in PAYMENT_TYPES]

    return [
        [entry.code, entry.name, entry.nif, entry.lei, entity_type(entry.type_code), iso(entry.registered_on)]
        for entry in sorted(active, key=lambda entry: entry.code)
    ]


def new_bank_rows(entries: tuple, banks: list) -> list:
    known = {bank["bank_code"] for bank in banks}
    missing = [entry for entry in entries if entry.deregistered_on is None and entry.type_code in CREDIT_TYPES and entry.code not in known]

    return [
        [entry.code, entry.name, entry.nif, entry.lei, entity_type(entry.type_code), entry.web]
        for entry in sorted(missing, key=lambda entry: entry.code)
    ]


def type_rows(entries: tuple) -> list:
    names = {entry.type_code: entry.type_name for entry in entries}
    return [[entity_type(code), code, names[code]] for code in ENTITY_TYPES if code in names]


def write_csv(path: Path, header: list, rows: list) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    path.write_text(buffer.getvalue(), encoding="utf-8", newline="")
    print(f"{path}: {len(rows)} filas")


def read_banks(path: Path) -> list:
    csv.field_size_limit(sys.maxsize)

    with path.open(encoding="utf-8", newline="") as file:
        rows = list(csv.DictReader(file))

    if not rows or "bank_code" not in rows[0]:
        raise ValueError(f"{path} no tiene la columna bank_code")

    return rows


def normalize(name: str) -> str:
    plain = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().casefold()
    return " ".join(re.sub(r"[^a-z0-9&]+", " ", plain).split())


def successor_status(code: str, index: Index) -> str:
    holder = active_entry(code, index)
    return f"{code} {holder.name}" if holder else f"{code} (sin entidad vigente)"


def report_banks(banks: list, index: Index) -> None:
    for bank in banks:
        code, name = bank["bank_code"], bank["bank_name"]
        entry = current_or_last(code, index)

        if entry is None:
            renumbering = RENUMBERED[code]
            print(f"  {code} {name!r}: pasó al {successor_status(renumbering.code, index)} el {iso(renumbering.until)} ({renumbering.source})")
        elif entry.deregistered_on:
            print(f"  {code} {name!r}: baja {iso(entry.deregistered_on)}, sucesora {successor_status(entry.successor_code, index)}")
        elif normalize(entry.name) != normalize(name):
            print(f"  {code} nombre distinto: {name!r} / registro {entry.name!r}")


def report_successors(historical: list, rows: list, index: Index) -> None:
    with_successor = [(entry, row) for entry, row in zip(historical, rows) if entry.successor_code]
    unknown = sorted({entry.successor_code for entry, _ in with_successor if entry.successor_code not in index.by_code})
    unresolved = [entry for entry, row in with_successor if not row[-1]]
    print(f"Sucesoras: {len(with_successor)} filas, {len(unresolved)} sin nombre resoluble")
    print(f"  Códigos sucesores que no aparecen en el registro: {', '.join(unknown) or 'ninguno'}")

    for entry in unresolved:
        print(f"  sin resolver: {entry.code} {iso(entry.deregistered_on)} -> {entry.successor_code}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera los CSV de bancos a partir del Registro de Entidades del Banco de España.")
    parser.add_argument("--xlsx", type=Path, help="Registro_Entidades.xlsx ya descargado")
    parser.add_argument("--bancos", type=Path, default=REPO / "bancos" / "bancos_es.csv")
    parser.add_argument("--out", type=Path, default=REPO / "bancos")
    parser.add_argument("--build", type=Path, default=REPO / "build")
    args = parser.parse_args()

    data = args.xlsx.read_bytes() if args.xlsx else fetch(SOURCE_URL)
    as_of, entries = read_register(data)
    index = build_index(entries)
    banks = read_banks(args.bancos)
    print(f"Registro de Entidades del BdE a {as_of}: {len(entries)} filas con código de 4 dígitos")

    historical = historical_entries(entries, index)
    historical_rows = [historical_row(entry, index) for entry in historical]
    args.out.mkdir(parents=True, exist_ok=True)
    args.build.mkdir(parents=True, exist_ok=True)
    write_csv(args.build / "bancos_enriquecimiento.csv", ENRICHMENT_HEADER, [enrichment_row(bank["bank_code"], index) for bank in banks])
    write_csv(args.build / "bancos_nuevos.csv", NEW_BANKS_HEADER, new_bank_rows(entries, banks))
    write_csv(args.out / "bancos_historicos_es.csv", HISTORICAL_HEADER, historical_rows)
    write_csv(args.out / "entidades_pago_es.csv", PAYMENTS_HEADER, payment_rows(entries))
    write_csv(args.out / "tipos_entidad_bde.csv", TYPES_HEADER, type_rows(entries))

    print("Revisión de bancos_es.csv:")
    report_banks(banks, index)
    report_successors(historical, historical_rows, index)


if __name__ == "__main__":
    main()
