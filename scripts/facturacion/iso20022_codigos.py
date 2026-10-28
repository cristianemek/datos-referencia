import argparse
import csv
import io
import sys
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
USER_AGENT = "Mozilla/5.0 (compatible; datos-referencia; +https://github.com/cristianemek/datos-referencia)"
ARCHIVE_URL = "https://www.iso20022.org/sites/default/files/media/file/ExternalCodeSets_XLSX.zip"
STATUSES = {"registered": True, "obsolete": False}

LISTS = (
    "ExternalServiceLevel1Code",
    "ExternalLocalInstrument1Code",
    "ExternalCategoryPurpose1Code",
    "ExternalPurpose1Code",
    "ExternalPaymentGroupStatus1Code",
    "ExternalPaymentTransactionStatus1Code",
    "ExternalStatusReason1Code",
    "ExternalReturnReason1Code",
    "ExternalReversalReason1Code",
    "ExternalCancellationReason1Code",
    "ExternalPaymentCancellationRejection1Code",
    "ExternalEntryStatus1Code",
    "ExternalBalanceType1Code",
)


def fetch(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})

    with urllib.request.urlopen(request, timeout=300) as response:
        return response.read()


def clean(text) -> str:
    return " ".join(str(text or "").replace("_x000D_", " ").split())


def open_workbook(archive: bytes):
    with zipfile.ZipFile(io.BytesIO(archive)) as bundle:
        names = [name for name in bundle.namelist() if name.lower().endswith(".xlsx")]

        if len(names) != 1:
            raise RuntimeError(f"Se esperaba un único XLSX en el ZIP de ISO 20022: {names}")

        return names[0], openpyxl.load_workbook(io.BytesIO(bundle.read(names[0])), read_only=True)


def publication_date(workbook) -> str:
    for row in workbook["Intro&History"].iter_rows(values_only=True):
        if row and row[0] == "Publication date":
            return clean(row[1])

    return "desconocida"


def code_rows(workbook) -> tuple:
    rows = workbook["AllCodeSets"].iter_rows(values_only=True)
    header = [clean(cell) for cell in next(rows)]
    column = {name: header.index(name) for name in ("Code Set", "Code Value", "Code Name", "Code Definition", "Status")}
    selected, obsolete = [], 0

    for row in rows:
        code_set = clean(row[column["Code Set"]])

        if code_set not in LISTS:
            continue

        status = clean(row[column["Status"]]).lower()

        if status not in STATUSES:
            raise RuntimeError(f"Estado desconocido en {code_set} {row[column['Code Value']]}: {status}")

        if not STATUSES[status]:
            obsolete += 1
            continue

        selected.append((
            code_set,
            clean(row[column["Code Value"]]),
            clean(row[column["Code Name"]]),
            clean(row[column["Code Definition"]]),
        ))

    return sorted(selected, key=lambda item: (LISTS.index(item[0]), item[1])), obsolete


def check(rows: list) -> None:
    keys = [(row[0], row[1]) for row in rows]
    missing = [name for name in LISTS if name not in {row[0] for row in rows}]

    if missing:
        raise RuntimeError(f"No se encuentran las listas {missing}")

    if len(keys) != len(set(keys)):
        raise RuntimeError("Hay códigos vigentes repetidos en una misma lista")

    if any(not code or not name for _, code, name, _ in rows):
        raise RuntimeError("Hay códigos sin valor o sin nombre")


def write_csv(path: Path, header: list, rows: list) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(buffer.getvalue(), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Códigos externos de ISO 20022 para pagos SEPA y conciliación")
    parser.add_argument("--output-dir", type=Path, default=REPO / "pagos")
    output = parser.parse_args().output_dir

    name, workbook = open_workbook(fetch(ARCHIVE_URL))
    published = publication_date(workbook)
    rows, obsolete = code_rows(workbook)
    check(rows)
    write_csv(output / "iso20022_codigos.csv", ["list", "code", "name", "definition"], rows)
    print(f"ISO 20022 {name} ({published}): {len(rows)} códigos en {len(LISTS)} listas, {obsolete} obsoletos descartados")


if __name__ == "__main__":
    sys.exit(main())
